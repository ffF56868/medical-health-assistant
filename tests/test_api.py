from datetime import UTC, datetime, timedelta
from io import BytesIO
from types import SimpleNamespace

from langchain_core.documents import Document
from sqlmodel import Session, create_engine, select

from app.database import create_db_and_tables
from app.document_parsers import ParsedDocument
from app.main import health_check
from app.models import (
    KnowledgeDocument,
    KnowledgeRebuildJob,
    RAGASAutoEvaluationRun,
)
from app.routers import ask as ask_router
from app.routers import evaluation as evaluation_router
from app.routers import knowledge as knowledge_router
from app.ragas_compat import enable_ragas_langchain_compatibility
from app.ragas_metrics import build_medical_answer_relevancy_metric
from app.seed_specialty_knowledge import (
    FOCUSED_KNOWLEDGE_DOCUMENTS,
    SPECIALTY_KNOWLEDGE_DOCUMENTS,
    seed_specialty_knowledge,
)
from app.routers import documents as documents_router
from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


class FakeVectorStore:
    def __init__(self, matches):
        self.matches = matches
        self.last_query = None
        self.last_k = None
        self.last_filter = None

    def similarity_search_with_relevance_scores(self, query, k, filter=None):
        self.last_query = query
        self.last_k = k
        self.last_filter = filter
        return self.matches


class FakeChatModel:
    def __init__(self, answer="这是测试模型回答。"):
        self.answer = answer
        self.call_count = 0

    def invoke(self, _messages):
        self.call_count += 1
        return SimpleNamespace(content=self.answer)

    def stream(self, _messages):
        self.call_count += 1
        midpoint = max(1, len(self.answer) // 2)
        yield SimpleNamespace(content=self.answer[:midpoint])
        yield SimpleNamespace(content=self.answer[midpoint:])


class FakeEvaluationVectorStore:
    def __init__(self):
        self.queries = []

    def similarity_search_with_relevance_scores(self, query, k):
        self.queries.append((query, k))
        case = next(
            item
            for item in evaluation_router.EVALUATION_CASES
            if item["question"] == query
        )
        return [
            (
                Document(
                    page_content="评测用资料",
                    metadata={
                        "name": case["expected_name"],
                        "type": case["expected_type"],
                    },
                ),
                0.9,
            )
        ]


class FakeComparisonVectorStore:
    def __init__(self, matches):
        self.matches = matches
        self.queries = []

    def similarity_search_with_relevance_scores(self, query, k):
        self.queries.append((query, k))
        return self.matches[:k]


def mock_current_knowledge_base(monkeypatch):
    monkeypatch.setattr(
        ask_router,
        "get_knowledge_status",
        lambda session: {"is_current": True},
    )


def build_ragas_question_result(question="RAGAS 测试问题"):
    return {
        "case_id": "ragas-test-case",
        "case_source": "测试题",
        "question": question,
        "expected_name": "测试资料",
        "expected_type": "document",
        "category": "测试",
        "answer": "这是依据测试资料生成的回答。",
        "processing_path": "rag-hybrid-rerank",
        "retrieved_count": 1,
        "context_titles": ["测试资料"],
        "retrieved_contexts": ["资料标题：测试资料\n资料内容：测试内容。"],
        "reference_available": True,
        "faithfulness": 0.9,
        "answer_relevancy": 0.8,
        "context_precision": 0.7,
        "context_recall": 0.6,
    }


def test_rag_answer_prompt_uses_a_direct_lookup_mode():
    messages = ask_router.build_rag_answer_prompt().format_messages(
        context="资料标题：焦虑障碍健康知识\n资料内容：焦虑可影响日常生活。",
        history="暂无历史对话",
        question="焦虑影响学习工作时应看哪份资料？",
        source_limitations="资料来自已审核来源。",
        reference_names="焦虑障碍健康知识",
        answer_mode=ask_router.build_answer_mode_instruction(
            "焦虑影响学习工作时应看哪份资料？"
        ),
    )

    system_message = messages[0].content
    assert "资料定位模式" in system_message
    assert "第一句必须直接回答" in system_message
    assert "焦虑障碍健康知识" in system_message


def test_rag_answer_mode_keeps_symptom_questions_in_the_safe_template():
    instruction = ask_router.build_answer_mode_instruction("我头痛又发热怎么办？")

    assert "症状咨询模式" in instruction
    assert "一、资料内容" in instruction


def test_vector_only_search_remains_available_without_a_database_session():
    vector_store = FakeVectorStore([])

    result = ask_router.search_knowledge(
        vector_store,
        "测试检索",
        "all",
        "all",
    )

    assert result == []
    assert vector_store.last_k == 8


def test_medical_ragas_metric_does_not_treat_safety_language_as_evasion():
    enable_ragas_langchain_compatibility()
    metric = build_medical_answer_relevancy_metric()

    instruction = metric.question_generation.instruction
    assert metric.name == "answer_relevancy"
    assert "医疗安全提示" in instruction
    assert "不得判为回避" in instruction


def test_ragas_task_defaults_to_ten_sampled_cases(client, monkeypatch):
    case = {
        "case_id": "ragas-task-case",
        "case_source": "测试题",
        "question": "测试 RAGAS 任务",
        "expected_name": "测试资料",
        "expected_type": "document",
        "category": "测试",
    }
    monkeypatch.setattr(evaluation_router, "EVALUATION_CASES", (case,))
    monkeypatch.setattr(
        evaluation_router,
        "get_knowledge_status",
        lambda session: {"is_current": True},
    )
    monkeypatch.setattr(
        evaluation_router,
        "run_ragas_evaluation_job",
        lambda _run_id: None,
    )

    response = client.post("/evaluation/ragas/tasks")

    assert response.status_code == 202
    data = response.json()
    assert data["status"] == "pending"
    assert data["sample_size"] == 10
    assert data["total_count"] == 0
    assert data["metrics"] is None

    history_response = client.get("/evaluation/ragas/tasks?limit=1")
    assert history_response.status_code == 200
    assert history_response.json()["total_count"] == 1
    assert history_response.json()["runs"][0]["id"] == data["id"]


def test_ragas_worker_persists_aggregate_and_question_scores(
    test_engine,
    monkeypatch,
):
    with Session(test_engine) as session:
        run = RAGASAutoEvaluationRun(sample_size=10)
        session.add(run)
        session.commit()
        session.refresh(run)
        run_id = run.id

    monkeypatch.setattr(evaluation_router, "engine", test_engine)
    monkeypatch.setattr(
        evaluation_router,
        "execute_ragas_evaluation",
        lambda _session, _run: [build_ragas_question_result()],
    )

    evaluation_router.run_ragas_evaluation_job(run_id)

    with Session(test_engine) as session:
        run = session.get(RAGASAutoEvaluationRun, run_id)
        assert run.status == "completed"
        assert run.completed_count == 1
        assert run.faithfulness == 0.9
        assert run.answer_relevancy == 0.8
        assert run.context_precision == 0.7
        assert run.context_recall == 0.6
        serialized = evaluation_router.serialize_ragas_run(run)
        assert serialized.results[0].question == "RAGAS 测试问题"
        assert serialized.results[0].context_titles == ["测试资料"]


def test_condition_can_be_created_and_listed(client):
    create_response = client.post(
        "/conditions",
        json={
            "name": "测试病症",
            "symptoms": "测试症状",
            "treatment": "测试处理建议",
            "source": "测试卫生机构",
            "source_url": "https://example.org/condition",
            "source_tier": "professional",
        },
    )
    assert create_response.status_code == 201

    list_response = client.get("/conditions", params={"keyword": "测试"})
    assert list_response.status_code == 200
    assert list_response.json()[0]["name"] == "测试病症"
    assert list_response.json()[0]["source_url"] == "https://example.org/condition"
    assert list_response.json()[0]["source_tier"] == "professional"
    assert list_response.json()[0]["updated_at"] is not None


def test_specialty_knowledge_seed_is_idempotent_and_preserves_existing_edits(
    test_engine,
):
    expected_document_count = len(
        SPECIALTY_KNOWLEDGE_DOCUMENTS + FOCUSED_KNOWLEDGE_DOCUMENTS
    )
    with Session(test_engine) as session:
        created_count, existing_count = seed_specialty_knowledge(session)
        assert created_count == expected_document_count
        assert existing_count == 0

        focused_document = session.exec(
            select(KnowledgeDocument).where(
                KnowledgeDocument.title == "哮喘：健康教育与就医警示"
            )
        ).one()
        assert focused_document.source_tier == "authority"
        assert focused_document.source_url is not None

        condition = session.exec(
            select(KnowledgeDocument).where(
                KnowledgeDocument.title == "呼吸内科常见病概览"
            )
        ).one()
        condition.content = "人工审核后的内容"
        session.add(condition)
        session.commit()

        created_count, existing_count = seed_specialty_knowledge(session)
        assert created_count == 0
        assert existing_count == expected_document_count

        preserved = session.exec(
            select(KnowledgeDocument).where(
                KnowledgeDocument.title == "呼吸内科常见病概览"
            )
        ).one()
        assert preserved.content == "人工审核后的内容"
        assert len(session.exec(select(KnowledgeDocument)).all()) == expected_document_count


def test_existing_sqlite_data_receives_metadata_columns_without_data_loss():
    old_database = create_engine("sqlite://")
    with old_database.begin() as connection:
        connection.exec_driver_sql(
            'CREATE TABLE "condition" ('
            'id INTEGER PRIMARY KEY, name VARCHAR(100), '
            'symptoms VARCHAR(5000), treatment VARCHAR(5000))'
        )
        connection.exec_driver_sql(
            "INSERT INTO \"condition\" (id, name, symptoms, treatment) "
            "VALUES (1, '旧资料', '旧症状', '旧建议')"
        )

    create_db_and_tables(old_database)

    with old_database.connect() as connection:
        columns = {
            row[1]
            for row in connection.exec_driver_sql('PRAGMA table_info("condition")')
        }
        migrated_row = connection.exec_driver_sql(
            'SELECT source, source_url, source_tier, updated_at '
            'FROM "condition" WHERE id = 1'
        ).one()

    assert {"source", "source_url", "source_tier", "updated_at"}.issubset(columns)
    assert migrated_row.source == "未标注来源"
    assert migrated_row.source_url is None
    assert migrated_row.source_tier == "unverified"
    assert migrated_row.updated_at is None


def test_existing_sqlite_evaluation_cases_receive_alternative_names_column():
    old_database = create_engine("sqlite://")
    with old_database.begin() as connection:
        connection.exec_driver_sql(
            'CREATE TABLE "ragevaluationcase" ('
            'id INTEGER PRIMARY KEY, question VARCHAR(1000), '
            'expected_name VARCHAR(200), expected_type VARCHAR(20), '
            'created_at DATETIME)'
        )
        connection.exec_driver_sql(
            "INSERT INTO \"ragevaluationcase\" "
            "(id, question, expected_name, expected_type, created_at) "
            "VALUES (1, '旧测试题', '旧资料', 'document', CURRENT_TIMESTAMP)"
        )

    create_db_and_tables(old_database)

    with old_database.connect() as connection:
        columns = {
            row[1]
            for row in connection.exec_driver_sql(
                'PRAGMA table_info("ragevaluationcase")'
            )
        }
        migrated_row = connection.exec_driver_sql(
            'SELECT category, alternative_names_json '
            'FROM "ragevaluationcase" WHERE id = 1'
        ).one()

    assert {"category", "alternative_names_json"}.issubset(columns)
    assert migrated_row.category == "自定义"
    assert migrated_row.alternative_names_json == "[]"


def test_existing_sqlite_evaluation_runs_receive_result_snapshot_column():
    old_database = create_engine("sqlite://")
    with old_database.begin() as connection:
        connection.exec_driver_sql(
            'CREATE TABLE "ragevaluationrun" ('
            'id INTEGER PRIMARY KEY, total_count INTEGER, passed_count INTEGER, '
            'pass_rate FLOAT, preset_count INTEGER, custom_count INTEGER, '
            'knowledge_document_count INTEGER, knowledge_hash VARCHAR(64), '
            'created_at DATETIME)'
        )
        connection.exec_driver_sql(
            "INSERT INTO \"ragevaluationrun\" "
            "(id, total_count, passed_count, pass_rate, preset_count, custom_count, "
            "knowledge_document_count, created_at) "
            "VALUES (1, 1, 1, 1, 1, 0, 1, CURRENT_TIMESTAMP)"
        )

    create_db_and_tables(old_database)

    with old_database.connect() as connection:
        migrated_row = connection.exec_driver_sql(
            'SELECT results_json FROM "ragevaluationrun" WHERE id = 1'
        ).one()

    assert migrated_row.results_json == "[]"


def test_existing_chat_messages_receive_response_metadata_column():
    old_database = create_engine("sqlite://")
    with old_database.begin() as connection:
        connection.exec_driver_sql(
            'CREATE TABLE "chatmessage" ('
            'id INTEGER PRIMARY KEY, conversation_id VARCHAR(100), '
            'role VARCHAR(20), content VARCHAR(10000), created_at DATETIME)'
        )
        connection.exec_driver_sql(
            "INSERT INTO \"chatmessage\" "
            "(id, conversation_id, role, content, created_at) "
            "VALUES (1, 'old-conversation', 'assistant', '旧回答', CURRENT_TIMESTAMP)"
        )

    create_db_and_tables(old_database)

    with old_database.connect() as connection:
        columns = {
            row[1]
            for row in connection.exec_driver_sql('PRAGMA table_info("chatmessage")')
        }
        migrated_row = connection.exec_driver_sql(
            'SELECT content, response_metadata_json FROM "chatmessage" WHERE id = 1'
        ).one()

    assert "response_metadata_json" in columns
    assert migrated_row.content == "旧回答"
    assert migrated_row.response_metadata_json == "{}"


def test_existing_sqlite_data_receives_the_review_log_table():
    old_database = create_engine("sqlite://")
    with old_database.begin() as connection:
        connection.exec_driver_sql(
            'CREATE TABLE "condition" ('
            'id INTEGER PRIMARY KEY, name VARCHAR(100), '
            'symptoms VARCHAR(5000), treatment VARCHAR(5000))'
        )

    create_db_and_tables(old_database)

    with old_database.connect() as connection:
        table_names = {
            row[0]
            for row in connection.exec_driver_sql(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }

    assert "knowledgereviewlog" in table_names


def test_existing_rebuild_jobs_receive_retry_reference_column():
    old_database = create_engine("sqlite://")
    with old_database.begin() as connection:
        connection.exec_driver_sql(
            'CREATE TABLE "knowledgerebuildjob" ('
            'id INTEGER PRIMARY KEY, status VARCHAR(20), '
            'document_count INTEGER, chunk_count INTEGER, '
            'error_message VARCHAR(2000), created_at DATETIME, '
            'started_at DATETIME, completed_at DATETIME)'
        )

    create_db_and_tables(old_database)

    with old_database.connect() as connection:
        columns = {
            row[1]
            for row in connection.exec_driver_sql(
                'PRAGMA table_info("knowledgerebuildjob")'
            )
        }

    assert "retry_of_job_id" in columns


def test_blank_knowledge_search_is_rejected(client):
    response = client.get("/knowledge/search", params={"q": " "})
    assert response.status_code == 422
    assert response.json()["detail"] == "搜索关键词不能为空"


def test_review_queue_lists_only_knowledge_that_needs_human_review(client):
    needs_review = client.post(
        "/conditions",
        json={
            "name": "待核验病症",
            "symptoms": "测试症状",
            "treatment": "测试处理建议",
        },
    )
    ready = client.post(
        "/drugs",
        json={
            "name": "已标注药物",
            "effects": "测试作用",
            "instructions": "测试说明",
            "source": "测试专业机构",
            "source_url": "https://example.org/drug",
            "source_tier": "professional",
        },
    )
    assert needs_review.status_code == 201
    assert ready.status_code == 201

    response = client.get("/knowledge/review-queue")

    assert response.status_code == 200
    data = response.json()
    assert data["total_count"] == 1
    assert data["results"][0]["title"] == "待核验病症"
    assert "可信度等级为待核实" in data["results"][0]["review_reasons"]
    assert "未标注具体资料来源" in data["results"][0]["review_reasons"]
    assert "未提供可访问的来源链接" in data["results"][0]["review_reasons"]


def test_source_url_must_be_an_http_address(client):
    response = client.post(
        "/conditions",
        json={
            "name": "错误链接病症",
            "symptoms": "测试症状",
            "treatment": "测试建议",
            "source": "测试机构",
            "source_url": "ftp://example.org/condition",
            "source_tier": "professional",
        },
    )

    assert response.status_code == 422
    assert "来源链接必须是有效的 http:// 或 https:// 地址" in response.text


def test_review_queue_batch_update_adds_verified_source_metadata(client):
    condition = client.post(
        "/conditions",
        json={
            "name": "批量审核病症",
            "symptoms": "测试症状",
            "treatment": "测试建议",
        },
    )
    drug = client.post(
        "/drugs",
        json={
            "name": "批量审核药物",
            "effects": "测试作用",
            "instructions": "测试说明",
        },
    )
    assert condition.status_code == 201
    assert drug.status_code == 201

    response = client.post(
        "/knowledge/review-queue/batch-update",
        json={
            "targets": [
                {"type": "condition", "record_id": condition.json()["id"]},
                {"type": "drug", "record_id": drug.json()["id"]},
            ],
            "source": "测试专业机构",
            "source_url": "https://example.org/reviewed-source",
            "source_tier": "professional",
        },
    )

    assert response.status_code == 200
    assert response.json()["updated_count"] == 2
    assert all(item["review_reasons"] == [] for item in response.json()["items"])
    assert client.get(f"/conditions/{condition.json()['id']}").json()["source_url"] == (
        "https://example.org/reviewed-source"
    )
    assert client.get(f"/drugs/{drug.json()['id']}").json()["source_tier"] == (
        "professional"
    )

    queue = client.get("/knowledge/review-queue")
    assert queue.status_code == 200
    assert queue.json()["total_count"] == 0

    history = client.get("/knowledge/review-logs")
    assert history.status_code == 200
    assert history.json()["total_count"] == 2
    assert {log["record_title"] for log in history.json()["logs"]} == {
        "批量审核病症",
        "批量审核药物",
    }
    assert all(
        log["source_url"] == "https://example.org/reviewed-source"
        for log in history.json()["logs"]
    )


def test_knowledge_versions_can_restore_a_previous_snapshot(client, monkeypatch):
    def fake_rebuild_vector_store(session):
        session.commit()
        return 1, 1

    monkeypatch.setattr(
        knowledge_router,
        "rebuild_vector_store",
        fake_rebuild_vector_store,
    )
    created = client.post(
        "/conditions",
        json={
            "name": "版本测试病症",
            "symptoms": "第一版症状",
            "treatment": "第一版建议",
        },
    )
    assert created.status_code == 201
    condition_id = created.json()["id"]

    first_rebuild = client.post("/knowledge/rebuild")
    assert first_rebuild.status_code == 200
    first_version_id = first_rebuild.json()["snapshot_id"]

    updated = client.put(
        f"/conditions/{condition_id}",
        json={
            "name": "版本测试病症",
            "symptoms": "第二版症状",
            "treatment": "第二版建议",
        },
    )
    assert updated.status_code == 200
    second_rebuild = client.post("/knowledge/rebuild")
    assert second_rebuild.status_code == 200
    assert second_rebuild.json()["snapshot_id"] != first_version_id

    restored = client.post(f"/knowledge/versions/{first_version_id}/restore")
    assert restored.status_code == 200
    assert restored.json()["restored_version_id"] == first_version_id
    assert restored.json()["backup_version_id"] == second_rebuild.json()["snapshot_id"]

    condition = client.get(f"/conditions/{condition_id}")
    assert condition.status_code == 200
    assert condition.json()["symptoms"] == "第一版症状"

    versions = client.get("/knowledge/versions")
    assert versions.status_code == 200
    restored_version = next(
        version
        for version in versions.json()["versions"]
        if version["id"] == first_version_id
    )
    assert restored_version["is_current"] is True


def test_async_rebuild_returns_job_and_can_query_completed_status(
    client,
    test_engine,
    monkeypatch,
):
    def fake_run_rebuild_job(job_id):
        with Session(test_engine) as session:
            job = session.get(KnowledgeRebuildJob, job_id)
            job.status = "completed"
            job.document_count = 2
            job.chunk_count = 5
            session.add(job)
            session.commit()

    monkeypatch.setattr(knowledge_router, "run_rebuild_job", fake_run_rebuild_job)

    response = client.post("/knowledge/rebuild/async")

    assert response.status_code == 202
    job_id = response.json()["id"]
    assert response.json()["status"] == "pending"

    status_response = client.get(f"/knowledge/rebuild/jobs/{job_id}")
    assert status_response.status_code == 200
    assert status_response.json()["document_count"] == 2
    assert status_response.json()["chunk_count"] == 5
    assert client.get("/knowledge/rebuild/jobs/active").json() is None


def test_async_rebuild_rejects_a_second_active_job(client, test_engine):
    with Session(test_engine) as session:
        active_job = KnowledgeRebuildJob(status="running")
        session.add(active_job)
        session.commit()
        session.refresh(active_job)

    response = client.post("/knowledge/rebuild/async")

    assert response.status_code == 409
    assert str(active_job.id) in response.json()["detail"]


def test_rebuild_job_records_failure(test_engine, monkeypatch):
    with Session(test_engine) as session:
        job = KnowledgeRebuildJob()
        session.add(job)
        session.commit()
        session.refresh(job)
        job_id = job.id

    monkeypatch.setattr(knowledge_router, "engine", test_engine)
    monkeypatch.setattr(
        knowledge_router,
        "rebuild_vector_store",
        lambda session: (_ for _ in ()).throw(RuntimeError("测试向量服务不可用")),
    )

    knowledge_router.run_rebuild_job(job_id)

    with Session(test_engine) as session:
        failed_job = session.get(KnowledgeRebuildJob, job_id)
        assert failed_job.status == "failed"
        assert failed_job.error_message == "测试向量服务不可用"
        assert failed_job.completed_at is not None


def test_rebuild_job_history_lists_recent_jobs(client, test_engine):
    with Session(test_engine) as session:
        session.add_all(
            [
                KnowledgeRebuildJob(
                    status="completed",
                    document_count=10,
                    chunk_count=20,
                ),
                KnowledgeRebuildJob(
                    status="failed",
                    error_message="测试失败",
                ),
            ]
        )
        session.commit()

    response = client.get("/knowledge/rebuild/jobs", params={"limit": 1})

    assert response.status_code == 200
    data = response.json()
    assert data["total_count"] == 2
    assert len(data["jobs"]) == 1
    assert data["jobs"][0]["status"] in {"completed", "failed"}


def test_failed_rebuild_job_can_create_a_retry(client, test_engine, monkeypatch):
    with Session(test_engine) as session:
        failed_job = KnowledgeRebuildJob(
            status="failed",
            error_message="第一次失败",
        )
        session.add(failed_job)
        session.commit()
        session.refresh(failed_job)
        failed_job_id = failed_job.id

    def fake_run_rebuild_job(job_id):
        with Session(test_engine) as session:
            job = session.get(KnowledgeRebuildJob, job_id)
            job.status = "completed"
            job.document_count = 3
            job.chunk_count = 7
            session.add(job)
            session.commit()

    monkeypatch.setattr(knowledge_router, "run_rebuild_job", fake_run_rebuild_job)

    response = client.post(f"/knowledge/rebuild/jobs/{failed_job_id}/retry")

    assert response.status_code == 202
    retry_job_id = response.json()["id"]
    assert response.json()["status"] == "pending"
    assert response.json()["retry_of_job_id"] == failed_job_id
    retry_status = client.get(f"/knowledge/rebuild/jobs/{retry_job_id}")
    assert retry_status.json()["status"] == "completed"
    assert retry_status.json()["chunk_count"] == 7


def test_completed_rebuild_job_cannot_be_retried(client, test_engine):
    with Session(test_engine) as session:
        job = KnowledgeRebuildJob(status="completed")
        session.add(job)
        session.commit()
        session.refresh(job)
        job_id = job.id

    response = client.post(f"/knowledge/rebuild/jobs/{job_id}/retry")

    assert response.status_code == 409
    assert response.json()["detail"] == "只有失败的知识库重建任务可以重试"


def test_stale_rebuild_job_is_marked_failed_and_can_be_seen(
    client,
    test_engine,
    monkeypatch,
):
    monkeypatch.setattr(knowledge_router, "REBUILD_JOB_TIMEOUT_SECONDS", 60)
    with Session(test_engine) as session:
        job = KnowledgeRebuildJob(
            status="running",
            started_at=datetime.now(UTC) - timedelta(seconds=61),
        )
        session.add(job)
        session.commit()
        session.refresh(job)
        job_id = job.id

    response = client.get(f"/knowledge/rebuild/jobs/{job_id}")

    assert response.status_code == 200
    assert response.json()["status"] == "failed"
    assert "可能因服务重启中断" in response.json()["error_message"]


def test_rag_evaluation_reports_hits_in_the_top_three(client, monkeypatch):
    case = {
        "case_id": "evaluation-hit",
        "case_source": "测试题",
        "question": "测试药物有什么作用？",
        "expected_name": "测试药物",
        "expected_type": "drug",
        "category": "用药信息",
    }
    vector_store = FakeEvaluationVectorStore()
    monkeypatch.setattr(evaluation_router, "EVALUATION_CASES", (case,))
    monkeypatch.setattr(
        evaluation_router,
        "get_knowledge_status",
        lambda session: {"is_current": True},
    )
    monkeypatch.setattr(
        evaluation_router,
        "get_vector_store",
        lambda: vector_store,
    )

    response = client.post("/evaluation/run")

    assert response.status_code == 200
    data = response.json()
    assert data["total_count"] == len(evaluation_router.EVALUATION_CASES)
    assert data["passed_count"] == data["total_count"]
    assert data["pass_rate"] == 1
    assert all(item["expected_rank"] == 1 for item in data["results"])
    assert all(k == 8 for _, k in vector_store.queries)
    assert data["quality_gate"]["status"] == "baseline"
    assert {metric["category"] for metric in data["category_metrics"]} == {"用药信息"}


def test_rag_evaluation_accepts_an_alternative_target_name(client, monkeypatch):
    case = {
        "case_id": "alternative-target-case",
        "case_source": "默认题",
        "question": "症状重叠时的相关资料是什么？",
        "expected_name": "普通感冒",
        "expected_type": "condition",
        "category": "症状相关",
        "alternative_names": ["过敏性鼻炎"],
    }
    vector_store = FakeComparisonVectorStore(
        [
            (
                Document(
                    page_content="过敏性鼻炎资料",
                    metadata={
                        "name": "过敏性鼻炎",
                        "type": "condition",
                        "record_id": 1,
                    },
                ),
                0.91,
            )
        ]
    )
    monkeypatch.setattr(evaluation_router, "EVALUATION_CASES", (case,))
    monkeypatch.setattr(
        evaluation_router,
        "get_knowledge_status",
        lambda session: {"is_current": True},
    )
    monkeypatch.setattr(evaluation_router, "get_vector_store", lambda: vector_store)
    monkeypatch.setattr(
        evaluation_router,
        "hybrid_search",
        lambda _session, store, question, **kwargs: store.similarity_search_with_relevance_scores(
            question,
            k=kwargs["vector_fetch_count"],
        ),
    )

    response = client.post("/evaluation/run")

    assert response.status_code == 200
    result = response.json()["results"][0]
    assert result["passed"] is True
    assert result["expected_rank"] == 1
    assert result["matched_name"] == "过敏性鼻炎"


def test_rag_evaluation_warns_when_a_previous_hit_regresses(client, monkeypatch):
    case = {
        "case_id": "regression-case",
        "case_source": "默认题",
        "question": "目标资料在哪里？",
        "expected_name": "目标资料",
        "expected_type": "document",
        "category": "检索回归测试",
    }
    vector_store = FakeComparisonVectorStore(
        [
            (
                Document(
                    page_content="目标资料内容",
                    metadata={"name": "目标资料", "type": "document", "record_id": 1},
                ),
                0.9,
            )
        ]
    )
    monkeypatch.setattr(evaluation_router, "EVALUATION_CASES", (case,))
    monkeypatch.setattr(
        evaluation_router,
        "get_knowledge_status",
        lambda session: {"is_current": True},
    )
    monkeypatch.setattr(evaluation_router, "get_vector_store", lambda: vector_store)

    first_response = client.post("/evaluation/run")
    assert first_response.status_code == 200
    assert first_response.json()["quality_gate"]["status"] == "baseline"

    vector_store.matches = [
        (
            Document(
                page_content="其他资料内容",
                metadata={"name": "其他资料", "type": "document", "record_id": 2},
            ),
            0.9,
        )
    ]
    second_response = client.post("/evaluation/run")

    assert second_response.status_code == 200
    quality_gate = second_response.json()["quality_gate"]
    assert quality_gate["status"] == "warning"
    assert quality_gate["regressed_questions"] == ["目标资料在哪里？"]


def test_rag_evaluation_reports_new_cases_separately_from_regressions(client, monkeypatch):
    baseline_case = {
        "case_id": "existing-case",
        "case_source": "默认题",
        "question": "已有测试题",
        "expected_name": "目标资料",
        "expected_type": "document",
        "category": "基础覆盖",
    }
    added_case = {
        "case_id": "added-case",
        "case_source": "默认题",
        "question": "新增测试题",
        "expected_name": "目标资料",
        "expected_type": "document",
        "category": "基础覆盖",
    }
    vector_store = FakeComparisonVectorStore(
        [
            (
                Document(
                    page_content="目标资料内容",
                    metadata={"name": "目标资料", "type": "document", "record_id": 1},
                ),
                0.9,
            )
        ]
    )
    monkeypatch.setattr(evaluation_router, "EVALUATION_CASES", (baseline_case,))
    monkeypatch.setattr(
        evaluation_router,
        "get_knowledge_status",
        lambda session: {"is_current": True},
    )
    monkeypatch.setattr(evaluation_router, "get_vector_store", lambda: vector_store)

    assert client.post("/evaluation/run").status_code == 200

    monkeypatch.setattr(
        evaluation_router,
        "EVALUATION_CASES",
        (baseline_case, added_case),
    )
    response = client.post("/evaluation/run")

    assert response.status_code == 200
    quality_gate = response.json()["quality_gate"]
    assert quality_gate["status"] == "expanded"
    assert quality_gate["new_questions"] == ["新增测试题"]
    assert quality_gate["regressed_questions"] == []


def test_rag_evaluation_flags_a_lower_rank_without_calling_it_a_failure(
    client,
    monkeypatch,
):
    case = {
        "case_id": "rank-regression-case",
        "case_source": "默认题",
        "question": "目标资料在哪里？",
        "expected_name": "目标资料",
        "expected_type": "document",
        "category": "检索回归测试",
    }
    vector_store = FakeComparisonVectorStore(
        [
            (
                Document(
                    page_content="目标资料内容",
                    metadata={"name": "目标资料", "type": "document", "record_id": 1},
                ),
                0.9,
            )
        ]
    )
    monkeypatch.setattr(evaluation_router, "EVALUATION_CASES", (case,))
    monkeypatch.setattr(
        evaluation_router,
        "get_knowledge_status",
        lambda session: {"is_current": True},
    )
    monkeypatch.setattr(evaluation_router, "get_vector_store", lambda: vector_store)

    assert client.post("/evaluation/run").status_code == 200

    vector_store.matches = [
        (
            Document(
                page_content="更相近资料",
                metadata={"name": "更相近资料", "type": "document", "record_id": 2},
            ),
            0.99,
        ),
        (
            Document(
                page_content="目标资料内容",
                metadata={"name": "目标资料", "type": "document", "record_id": 1},
            ),
            0.9,
        ),
    ]
    monkeypatch.setattr(
        evaluation_router,
        "hybrid_search",
        lambda _session, store, question, **kwargs: store.similarity_search_with_relevance_scores(
            question,
            k=kwargs["vector_fetch_count"],
        ),
    )
    response = client.post("/evaluation/run")

    assert response.status_code == 200
    data = response.json()
    assert data["results"][0]["passed"] is True
    assert data["results"][0]["expected_rank"] == 2
    assert data["quality_gate"]["status"] == "attention"
    assert data["quality_gate"]["rank_regressed_questions"] == ["目标资料在哪里？"]


def test_retrieval_comparison_reports_deduplication_improvement(client, monkeypatch):
    case = {
        "case_id": "deduplication-case",
        "case_source": "默认题",
        "question": "目标资料在哪里？",
        "expected_name": "目标资料",
        "expected_type": "document",
    }
    vector_store = FakeComparisonVectorStore(
        [
            (
                Document(
                    page_content="同一资料的第一个切块",
                    metadata={"name": "重复资料", "type": "document", "record_id": 1},
                ),
                0.99,
            ),
            (
                Document(
                    page_content="同一资料的第二个切块",
                    metadata={"name": "重复资料", "type": "document", "record_id": 1},
                ),
                0.98,
            ),
            (
                Document(
                    page_content="另一份资料",
                    metadata={"name": "另一份资料", "type": "document", "record_id": 2},
                ),
                0.97,
            ),
            (
                Document(
                    page_content="目标资料内容",
                    metadata={"name": "目标资料", "type": "document", "record_id": 3},
                ),
                0.96,
            ),
        ]
    )
    monkeypatch.setattr(evaluation_router, "EVALUATION_CASES", (case,))
    monkeypatch.setattr(
        evaluation_router,
        "get_knowledge_status",
        lambda session: {"is_current": True},
    )
    monkeypatch.setattr(
        evaluation_router,
        "get_vector_store",
        lambda: vector_store,
    )
    monkeypatch.setattr(
        evaluation_router,
        "hybrid_search",
        lambda _session, store, question, **kwargs: store.similarity_search_with_relevance_scores(
            question,
            k=kwargs["vector_fetch_count"],
        ),
    )

    response = client.post("/evaluation/compare")

    assert response.status_code == 200
    data = response.json()
    assert data["baseline"]["passed_count"] == 0
    assert data["current"]["passed_count"] == 1
    assert data["pass_rate_delta"] == 1
    assert data["improved_count"] == 1
    assert data["regressed_count"] == 0
    assert data["results"][0]["baseline"]["expected_rank"] is None
    assert data["results"][0]["current"]["expected_rank"] == 3
    assert data["results"][0]["change"] == "improved"
    assert [k for _, k in vector_store.queries] == [3, 8]


def test_retrieval_diagnosis_explains_a_lower_rank_target(client, monkeypatch):
    case = {
        "case_id": "diagnosis-case",
        "case_source": "默认题",
        "question": "目标资料在哪里？",
        "expected_name": "目标资料",
        "expected_type": "document",
    }
    vector_store = FakeComparisonVectorStore(
        [
            (
                Document(
                    page_content="更相近资料",
                    metadata={"name": "更相近资料", "type": "document", "record_id": 1},
                ),
                0.92,
            ),
            (
                Document(
                    page_content="目标资料内容",
                    metadata={"name": "目标资料", "type": "document", "record_id": 2},
                ),
                0.81,
            ),
        ]
    )
    monkeypatch.setattr(evaluation_router, "EVALUATION_CASES", (case,))
    monkeypatch.setattr(
        evaluation_router,
        "get_knowledge_status",
        lambda session: {"is_current": True},
    )
    monkeypatch.setattr(
        evaluation_router,
        "get_vector_store",
        lambda: vector_store,
    )
    monkeypatch.setattr(
        evaluation_router,
        "hybrid_search",
        lambda _session, store, question, **kwargs: store.similarity_search_with_relevance_scores(
            question,
            k=kwargs["vector_fetch_count"],
        ),
    )

    response = client.post("/evaluation/diagnose")

    assert response.status_code == 200
    data = response.json()
    assert data["healthy_count"] == 0
    assert data["attention_count"] == 1
    assert data["failed_count"] == 0
    assert data["results"][0]["expected_rank"] == 2
    assert data["results"][0]["diagnostic_level"] == "attention"
    assert data["results"][0]["candidates"][0]["name"] == "更相近资料"


def test_retrieval_diagnosis_explains_a_missing_target(client, monkeypatch):
    case = {
        "case_id": "missing-diagnosis-case",
        "case_source": "默认题",
        "question": "目标资料在哪里？",
        "expected_name": "目标资料",
        "expected_type": "document",
    }
    vector_store = FakeComparisonVectorStore(
        [
            (
                Document(
                    page_content="不相关资料",
                    metadata={"name": "不相关资料", "type": "document", "record_id": 1},
                ),
                0.91,
            )
        ]
    )
    monkeypatch.setattr(evaluation_router, "EVALUATION_CASES", (case,))
    monkeypatch.setattr(
        evaluation_router,
        "get_knowledge_status",
        lambda session: {"is_current": True},
    )
    monkeypatch.setattr(
        evaluation_router,
        "get_vector_store",
        lambda: vector_store,
    )

    response = client.post("/evaluation/diagnose")

    assert response.status_code == 200
    data = response.json()
    assert data["failed_count"] == 1
    assert data["results"][0]["diagnostic_level"] == "failed"
    assert "没有进入前 8 个向量候选" in data["results"][0]["diagnostic"]


def test_rag_evaluation_saves_and_lists_history(client, monkeypatch):
    vector_store = FakeEvaluationVectorStore()
    monkeypatch.setattr(
        evaluation_router,
        "get_knowledge_status",
        lambda session: {
            "is_current": True,
            "document_count": 13,
        },
    )
    monkeypatch.setattr(
        evaluation_router,
        "get_vector_store",
        lambda: vector_store,
    )

    run_response = client.post("/evaluation/run")

    assert run_response.status_code == 200
    run_data = run_response.json()
    assert run_data["history_id"] > 0

    history_response = client.get("/evaluation/history")

    assert history_response.status_code == 200
    history_data = history_response.json()
    assert history_data["total_count"] == 1
    assert history_data["runs"][0]["id"] == run_data["history_id"]
    assert history_data["runs"][0]["passed_count"] == len(
        evaluation_router.EVALUATION_CASES
    )
    assert history_data["runs"][0]["knowledge_document_count"] == 13


def test_evaluation_history_limit_is_validated(client):
    response = client.get("/evaluation/history", params={"limit": 0})

    assert response.status_code == 422


def test_custom_evaluation_cases_can_be_created_listed_and_deleted(client):
    created = client.post(
        "/evaluation/cases",
        json={
            "question": "布洛芬有什么作用？",
            "expected_name": "布洛芬",
            "expected_type": "drug",
            "category": "用药信息",
            "alternative_names": ["布洛芬缓释胶囊", "布洛芬"],
        },
    )

    assert created.status_code == 201
    case_id = created.json()["id"]
    assert created.json()["expected_type"] == "drug"
    assert created.json()["category"] == "用药信息"
    assert created.json()["alternative_names"] == ["布洛芬缓释胶囊", "布洛芬"]

    listed = client.get("/evaluation/cases")
    assert listed.status_code == 200
    assert listed.json()["total_count"] == 1
    assert listed.json()["cases"][0]["id"] == case_id
    assert listed.json()["cases"][0]["category"] == "用药信息"
    assert listed.json()["cases"][0]["alternative_names"] == ["布洛芬缓释胶囊", "布洛芬"]

    deleted = client.delete(f"/evaluation/cases/{case_id}")
    assert deleted.status_code == 204
    assert client.get("/evaluation/cases").json()["cases"] == []


def test_custom_evaluation_case_rejects_unknown_knowledge_type(client):
    response = client.post(
        "/evaluation/cases",
        json={
            "question": "测试问题",
            "expected_name": "测试资料",
            "expected_type": "unknown",
        },
    )

    assert response.status_code == 422


def test_condition_can_be_updated_and_deleted(client):
    created = client.post(
        "/conditions",
        json={
            "name": "待编辑病症",
            "symptoms": "旧症状",
            "treatment": "旧建议",
        },
    )
    condition_id = created.json()["id"]

    updated = client.put(
        f"/conditions/{condition_id}",
        json={
            "name": "已编辑病症",
            "symptoms": "新症状",
            "treatment": "新建议",
        },
    )
    assert updated.status_code == 200
    assert updated.json()["name"] == "已编辑病症"

    deleted = client.delete(f"/conditions/{condition_id}")
    assert deleted.status_code == 204
    assert client.get(f"/conditions/{condition_id}").status_code == 404


def test_text_document_can_be_uploaded(client):
    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "睡眠提示.txt",
                "保持规律作息，睡前避免长时间使用电子设备。".encode("utf-8"),
                "text/plain",
            )
        },
    )
    assert response.status_code == 201
    assert response.json()["title"] == "睡眠提示"


def test_supported_documents_can_be_uploaded_in_a_batch_with_per_file_errors(client):
    word_document = WordDocument()
    word_document.add_paragraph("Word 中的健康知识")
    word_buffer = BytesIO()
    word_document.save(word_buffer)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "用药表"
    worksheet.append(["药物", "注意事项"])
    worksheet.append(["布洛芬", "按说明书使用"])
    excel_buffer = BytesIO()
    workbook.save(excel_buffer)

    response = client.post(
        "/documents/upload-batch",
        files=[
            (
                "files",
                (
                    "呼吸提示.md",
                    "保持室内通风，出现呼吸困难及时就医。".encode("utf-8"),
                    "text/markdown",
                ),
            ),
            (
                "files",
                (
                    "健康知识.docx",
                    word_buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            ),
            (
                "files",
                (
                    "用药资料.xlsx",
                    excel_buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            (
                "files",
                ("不支持.zip", b"not-supported", "application/zip"),
            ),
        ],
    )

    assert response.status_code == 201
    data = response.json()
    assert data["created_count"] == 3
    assert data["created_document_count"] == 3
    assert data["failed_count"] == 1
    assert data["items"][0]["title"] == "呼吸提示"
    assert data["errors"][0]["filename"] == "不支持.zip"
    assert "只支持上传" in data["errors"][0]["detail"]


def test_image_only_pdf_is_rejected_with_an_ocr_hint(client):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    image_only_buffer = BytesIO()
    writer.write(image_only_buffer)

    response = client.post(
        "/documents/upload",
        files={"file": ("扫描件.pdf", image_only_buffer.getvalue(), "application/pdf")},
    )

    assert response.status_code == 400
    assert "OCR" in response.json()["detail"]


def test_text_pdf_is_imported_with_its_page_number(client):
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): font_reference}
            )
        }
    )
    text_stream = DecodedStreamObject()
    text_stream.set_data(b"BT /F1 12 Tf 20 100 Td (Health PDF text) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(text_stream)
    pdf_buffer = BytesIO()
    writer.write(pdf_buffer)

    response = client.post(
        "/documents/upload",
        files={"file": ("用药指南.pdf", pdf_buffer.getvalue(), "application/pdf")},
    )

    assert response.status_code == 201
    assert response.json()["title"] == "用药指南（第1页）"
    assert response.json()["page_number"] == 1
    assert "Health PDF text" in response.json()["content"]


def test_web_page_can_be_imported(client, monkeypatch):
    monkeypatch.setattr(
        documents_router,
        "parse_web_page",
        lambda _url, _title: ParsedDocument(
            title="测试健康网页",
            content="网页中的健康资料正文",
            source="网页导入：https://example.org/health",
            source_url="https://example.org/health",
        ),
    )

    response = client.post(
        "/documents/import-web",
        json={"url": "https://example.org/health"},
    )

    assert response.status_code == 201
    assert response.json()["title"] == "测试健康网页"
    assert response.json()["source_url"] == "https://example.org/health"


def test_conversation_can_be_listed_read_and_deleted(client):
    ask_response = client.post(
        "/ask",
        json={
            "conversation_id": "test-conversation-history",
            "question": "我出现持续胸痛怎么办？",
        },
    )
    assert ask_response.status_code == 200

    summaries = client.get("/conversations")
    assert summaries.status_code == 200
    assert summaries.json()[0]["conversation_id"] == "test-conversation-history"
    assert summaries.json()[0]["message_count"] == 2

    messages = client.get("/conversations/test-conversation-history/messages")
    assert messages.status_code == 200
    assert len(messages.json()) == 2

    deleted = client.delete("/conversations/test-conversation-history/messages")
    assert deleted.status_code == 204
    assert client.get("/conversations").json() == []


def test_conversation_can_be_exported_as_markdown(client):
    conversation_id = "test-conversation-export"
    ask_response = client.post(
        "/ask",
        json={
            "conversation_id": conversation_id,
            "question": "我出现持续胸痛怎么办？",
        },
    )
    assert ask_response.status_code == 200

    export_response = client.get(f"/conversations/{conversation_id}/export")

    assert export_response.status_code == 200
    assert export_response.headers["content-type"].startswith("text/markdown")
    assert "attachment; filename=medical-health-conversation.md" in export_response.headers[
        "content-disposition"
    ]
    assert "我出现持续胸痛怎么办？" in export_response.text
    assert "需要及时就医的警示信号" in export_response.text
    assert "内容仅供健康信息参考，不代替医生诊断或处方。" in export_response.text

    missing_response = client.get("/conversations/missing-conversation/export")
    assert missing_response.status_code == 404


def test_urgent_warning_bypasses_model_and_can_receive_feedback(client):
    ask_response = client.post(
        "/ask",
        json={
            "conversation_id": "test-urgent-feedback",
            "question": "我现在感觉呼吸困难，怎么办？",
        },
    )
    assert ask_response.status_code == 200
    ask_data = ask_response.json()
    assert ask_data["source"] == "safety-keyword-guard"
    assert "立即寻求紧急医疗帮助" in ask_data["answer"]

    feedback_response = client.post(
        "/feedback",
        json={
            "assistant_message_id": ask_data["assistant_message_id"],
            "helpful": False,
            "reason": "资料不足，没有回答这个问题",
        },
    )
    assert feedback_response.status_code == 200

    suggestions_response = client.get("/feedback/improvement-suggestions")
    assert suggestions_response.status_code == 200
    suggestions = suggestions_response.json()
    assert suggestions["not_helpful_count"] == 1
    assert suggestions["repeated_questions"][0]["text"] == "我现在感觉呼吸困难，怎么办？"


def test_health_check_reports_database_and_knowledge_state(test_engine, monkeypatch):
    expected_status = {
        "is_current": True,
        "document_count": 0,
        "chunk_count": 0,
    }
    monkeypatch.setattr(
        "app.main.get_knowledge_status",
        lambda session: expected_status,
    )
    with Session(test_engine) as session:
        response = health_check(session)

    assert response.status == "ok"
    assert response.database == "connected"
    assert response.knowledge_base_current is True


def test_rag_keeps_the_best_chunk_and_returns_structured_references(
    client,
    monkeypatch,
):
    mock_current_knowledge_base(monkeypatch)
    vector_store = FakeVectorStore(
        [
            (
                Document(
                    page_content="睡眠资料的低分切块",
                    metadata={
                        "type": "document",
                        "record_id": 1,
                        "name": "睡眠健康提示",
                        "source": "测试文档",
                        "source_url": "https://example.org/sleep",
                        "chunk_index": 0,
                    },
                ),
                0.5,
            ),
            (
                Document(
                    page_content="睡眠资料的高分切块",
                    metadata={
                        "type": "document",
                        "record_id": 1,
                        "name": "睡眠健康提示",
                        "source": "测试文档",
                        "source_url": "https://example.org/sleep",
                        "chunk_index": 1,
                    },
                ),
                0.9,
            ),
            (
                Document(
                    page_content="布洛芬可用于缓解发热和疼痛。",
                    metadata={
                        "type": "drug",
                        "record_id": 2,
                        "name": "布洛芬",
                        "chunk_index": 0,
                    },
                ),
                0.7,
            ),
        ]
    )
    chat_model = FakeChatModel()
    monkeypatch.setattr(ask_router, "get_vector_store", lambda: vector_store)
    monkeypatch.setattr(ask_router, "get_chat_model", lambda: chat_model)

    response = client.post(
        "/ask",
        json={"conversation_id": "test-rag-references", "question": "睡眠问题"},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["answer"] == "这是测试模型回答。"
    assert data["source"] == "vector-rerank-retrieval-openai-generation"
    assert vector_store.last_k == 8
    assert vector_store.last_filter is None
    assert chat_model.call_count == 1
    assert [item["name"] for item in data["references"]] == [
        "睡眠健康提示",
        "布洛芬",
    ]
    assert data["references"][0]["relevance_score"] == 0.9
    assert data["references"][0]["record_id"] == 1
    assert data["references"][1]["record_id"] == 2
    assert "高分切块" in data["references"][0]["excerpt"]
    assert data["references"][0]["source"] == "测试文档"
    assert data["references"][0]["source_url"] == "https://example.org/sleep"
    assert data["references"][0]["source_tier"] == "unverified"
    assert data["references"][0]["needs_review"] is True
    assert data["processing_path"] == "rag-vector-rerank"
    assert data["retrieval_scope"] == "all"
    assert data["source_filter"] == "all"
    assert data["retrieved_count"] == 2
    assert data["latency_ms"] >= 0

    messages_response = client.get("/conversations/test-rag-references/messages")
    assert messages_response.status_code == 200
    assistant_metadata = messages_response.json()[1]["response_metadata"]
    assert assistant_metadata["processing_path"] == "rag-vector-rerank"
    assert assistant_metadata["retrieval_scope"] == "all"
    assert assistant_metadata["source_filter"] == "all"
    assert assistant_metadata["references"][0]["record_id"] == 1


def test_rag_returns_no_match_without_calling_the_chat_model(client, monkeypatch):
    mock_current_knowledge_base(monkeypatch)
    vector_store = FakeVectorStore(
        [
            (
                Document(
                    page_content="不相关资料",
                    metadata={"type": "document", "record_id": 1, "name": "资料"},
                ),
                0.19,
            )
        ]
    )
    monkeypatch.setattr(ask_router, "get_vector_store", lambda: vector_store)
    monkeypatch.setattr(
        ask_router,
        "get_chat_model",
        lambda: (_ for _ in ()).throw(AssertionError("不应调用模型")),
    )

    response = client.post(
        "/ask",
        json={"conversation_id": "test-rag-no-match", "question": "无关问题"},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["source"] == "milvus-vector-search:no-match"
    assert data["references"] == []


def test_rag_prefers_a_document_with_a_direct_title_match(client, monkeypatch):
    mock_current_knowledge_base(monkeypatch)
    vector_store = FakeVectorStore(
        [
            (
                Document(
                    page_content="哮喘资料内容。",
                    metadata={
                        "type": "document",
                        "record_id": 1,
                        "name": "哮喘：健康教育与就医警示",
                    },
                ),
                0.55,
            ),
            (
                Document(
                    page_content="高血压资料内容。",
                    metadata={
                        "type": "document",
                        "record_id": 2,
                        "name": "高血压：健康教育与就医警示",
                    },
                ),
                0.53,
            ),
        ]
    )
    chat_model = FakeChatModel()
    monkeypatch.setattr(ask_router, "get_vector_store", lambda: vector_store)
    monkeypatch.setattr(ask_router, "get_chat_model", lambda: chat_model)

    response = client.post(
        "/ask",
        json={
            "conversation_id": "test-rag-title-match",
            "question": "哮喘有哪些常见表现？",
            "knowledge_type": "document",
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["retrieved_count"] == 1
    assert [item["name"] for item in data["references"]] == [
        "哮喘：健康教育与就医警示"
    ]


def test_rag_limits_vector_search_to_the_selected_knowledge_type(client, monkeypatch):
    mock_current_knowledge_base(monkeypatch)
    vector_store = FakeVectorStore(
        [
            (
                Document(
                    page_content="氯雷他定用于缓解过敏相关不适。",
                    metadata={"type": "drug", "record_id": 3, "name": "氯雷他定"},
                ),
                0.9,
            )
        ]
    )
    monkeypatch.setattr(ask_router, "get_vector_store", lambda: vector_store)
    monkeypatch.setattr(ask_router, "get_chat_model", FakeChatModel)

    response = client.post(
        "/ask",
        json={
            "conversation_id": "test-rag-drug-scope",
            "question": "相关药物资料是什么？",
            "knowledge_type": "drug",
        },
    )

    assert response.status_code == 200
    assert vector_store.last_filter == {"type": "drug"}
    assert response.json()["retrieval_scope"] == "drug"


def test_rag_can_limit_vector_search_to_reviewed_sources(client, monkeypatch):
    mock_current_knowledge_base(monkeypatch)
    vector_store = FakeVectorStore(
        [
            (
                Document(
                    page_content="资料内容。",
                    metadata={
                        "type": "drug",
                        "record_id": 3,
                        "name": "测试药物",
                        "source": "测试专业机构",
                        "source_tier": "professional",
                        "needs_review": False,
                    },
                ),
                0.9,
            )
        ]
    )
    monkeypatch.setattr(ask_router, "get_vector_store", lambda: vector_store)
    monkeypatch.setattr(ask_router, "get_chat_model", FakeChatModel)

    response = client.post(
        "/ask",
        json={
            "conversation_id": "test-rag-reviewed-scope",
            "question": "测试药物资料是什么？",
            "knowledge_type": "drug",
            "source_filter": "reviewed",
        },
    )

    assert response.status_code == 200
    assert vector_store.last_filter == {
        "$and": [{"type": "drug"}, {"needs_review": False}]
    }
    assert response.json()["source_filter"] == "reviewed"


def test_rag_rejects_an_unknown_knowledge_type(client):
    response = client.post(
        "/ask",
        json={
            "conversation_id": "test-rag-invalid-scope",
            "question": "布洛芬有什么作用？",
            "knowledge_type": "unknown",
        },
    )

    assert response.status_code == 422
    assert "检索范围必须是 all、condition、drug 或 document" in response.text


def test_rag_rejects_an_unknown_source_filter(client):
    response = client.post(
        "/ask",
        json={
            "conversation_id": "test-rag-invalid-source-filter",
            "question": "布洛芬有什么作用？",
            "source_filter": "trusted",
        },
    )

    assert response.status_code == 422
    assert "资料可信度筛选必须是 all 或 reviewed" in response.text


def test_rag_rejects_requests_when_the_knowledge_base_is_outdated(
    client,
    monkeypatch,
):
    monkeypatch.setattr(
        ask_router,
        "get_knowledge_status",
        lambda session: {"is_current": False},
    )

    response = client.post(
        "/ask",
        json={"conversation_id": "test-rag-stale", "question": "布洛芬的作用"},
    )

    assert response.status_code == 409
    assert response.json()["detail"] == "知识库已过期，请先执行 POST /knowledge/rebuild"


def test_rag_streams_tokens_and_saves_the_completed_answer(client, monkeypatch):
    mock_current_knowledge_base(monkeypatch)
    vector_store = FakeVectorStore(
        [
            (
                Document(
                    page_content="布洛芬可用于缓解发热和疼痛。",
                    metadata={"type": "drug", "record_id": 2, "name": "布洛芬"},
                ),
                0.9,
            )
        ]
    )
    chat_model = FakeChatModel(answer="这是分段返回的测试回答。")
    monkeypatch.setattr(ask_router, "get_vector_store", lambda: vector_store)
    monkeypatch.setattr(ask_router, "get_chat_model", lambda: chat_model)

    response = client.post(
        "/ask/stream",
        json={
            "conversation_id": "test-streaming-answer",
            "question": "布洛芬有什么作用？",
            "knowledge_type": "drug",
            "source_filter": "reviewed",
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/event-stream")
    assert "event: metadata" in response.text
    assert response.text.count("event: token") == 2
    assert '"text": "这是分段返回"' in response.text
    assert '"text": "的测试回答。"' in response.text
    assert '"processing_path": "rag-vector-rerank"' in response.text
    assert '"retrieval_scope": "drug"' in response.text
    assert '"source_filter": "reviewed"' in response.text
    assert '"retrieved_count": 1' in response.text
    assert "event: done" in response.text
    assert chat_model.call_count == 1
    assert vector_store.last_filter == {
        "$and": [{"type": "drug"}, {"needs_review": False}]
    }

    messages = client.get("/conversations/test-streaming-answer/messages")
    assert [message["content"] for message in messages.json()] == [
        "布洛芬有什么作用？",
        "这是分段返回的测试回答。",
    ]
    assistant_metadata = messages.json()[1]["response_metadata"]
    assert assistant_metadata["processing_path"] == "rag-vector-rerank"
    assert assistant_metadata["retrieval_scope"] == "drug"
    assert assistant_metadata["source_filter"] == "reviewed"
    assert assistant_metadata["references"][0]["name"] == "布洛芬"
