"""Convert supported knowledge sources into plain text records.

This module deliberately only parses input. Cleaning and chunking are handled
in later knowledge-base stages, so the original meaning stays traceable.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from bs4 import BeautifulSoup
from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader

# OCR dependencies - optional, only used for scanned PDFs
try:
    import pytesseract
    from pdf2image import convert_from_bytes
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False


ALLOWED_SUFFIXES = {".md", ".txt", ".pdf", ".docx", ".xlsx", ".xlsm"}
MAX_DOCUMENT_BYTES = 5_000_000
MAX_BATCH_BYTES = 20_000_000
MAX_PARSED_TEXT_CHARS = 20_000
MAX_WEB_BYTES = 1_000_000


@dataclass(frozen=True)
class ParsedDocument:
    title: str
    content: str
    source: str
    source_url: str | None = None
    page_number: int | None = None


def get_document_title(filename: str) -> str:
    title = Path(filename).stem.strip()
    if not title:
        raise ValueError("文件名不能为空")
    return title[:200]


def validate_upload_file(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise ValueError(
            "只支持上传 .md、.txt、.pdf、.docx、.xlsx 或 .xlsm 文件"
        )
    return suffix


def normalize_text(text: str, label: str) -> str:
    normalized = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not normalized:
        raise ValueError(f"{label}中没有可读取的文字内容")
    if len(normalized) > MAX_PARSED_TEXT_CHARS:
        raise ValueError(
            f"{label}解析后的文字超过 {MAX_PARSED_TEXT_CHARS:,} 个字符"
        )
    return normalized


def parse_text_file(raw_content: bytes, filename: str) -> list[ParsedDocument]:
    try:
        content = raw_content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("文本文件必须使用 UTF-8 编码保存") from error
    return [
        ParsedDocument(
            title=get_document_title(filename),
            content=normalize_text(content, filename),
            source=f"上传文件：{filename}",
        )
    ]


def _ocr_page(page_image) -> str:
    """对单页图片进行 OCR 识别"""
    if not OCR_AVAILABLE:
        return ""
    try:
        # 使用中文+英文识别
        text = pytesseract.image_to_string(page_image, lang='chi_sim+eng')
        return text.strip()
    except Exception:
        return ""


def _is_scan_page(text: str) -> bool:
    """判断是否为扫描件页面（文本为空或极少）"""
    return len(text.strip()) < 50


def _ends_with_punctuation(text: str) -> bool:
    """检查文本是否以结束标点结尾"""
    if not text.strip():
        return False
    last_char = text.strip()[-1]
    return last_char in '。！？；.!?;'


def parse_pdf_file(raw_content: bytes, filename: str) -> list[ParsedDocument]:
    try:
        reader = PdfReader(BytesIO(raw_content))
    except Exception as error:
        raise ValueError("PDF 文件无法读取，可能已损坏或受密码保护") from error

    base_title = get_document_title(filename)
    parsed_documents: list[ParsedDocument] = []
    
    # 用于跨页段落拼接
    carry_over_text = ""
    carry_over_page = 0
    
    # 检查是否有扫描件页面，如果有且 OCR 可用，预先转换所有页面图片
    needs_ocr = False
    page_texts = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        page_texts.append(page_text)
        if _is_scan_page(page_text):
            needs_ocr = True
    
    # 如果有扫描件页面且 OCR 可用，转换图片
    ocr_images = []
    if needs_ocr and OCR_AVAILABLE:
        try:
            ocr_images = convert_from_bytes(raw_content)
        except Exception:
            # OCR 转换失败，继续用原文本
            ocr_images = []
    
    for page_number, page_text in enumerate(page_texts, start=1):
        # 如果是扫描件页面且有 OCR 图片，进行 OCR
        if _is_scan_page(page_text) and ocr_images and page_number <= len(ocr_images):
            page_text = _ocr_page(ocr_images[page_number - 1])
        
        if not page_text.strip():
            continue
        
        # 跨页段落拼接逻辑：如果上一页有未结束的段落，拼接到当前页开头
        if carry_over_text:
            page_text = carry_over_text + page_text
            carry_over_text = ""
        
        try:
            normalized_content = normalize_text(page_text, f"PDF 第 {page_number} 页")
        except ValueError:
            # 内容过长或为空，跳过
            continue
            
        parsed_documents.append(
            ParsedDocument(
                title=f"{base_title}（第{page_number}页）"[:200],
                content=normalized_content,
                source=f"上传 PDF：{filename}",
                page_number=page_number,
            )
        )
        
        # 检查当前页末尾是否需要延续到下一页（用于下一页的拼接）
        # 当前页内容已经保存，不影响当前页的完整性
        if not _ends_with_punctuation(page_text) and page_number < len(page_texts):
            # 取最后一行作为可能的跨页段落
            lines = page_text.rstrip().split('\n')
            if lines:
                last_line = lines[-1]
                # 如果最后一行较短且不以标点结尾，认为是跨页段落
                if len(last_line) < 100 and not _ends_with_punctuation(last_line):
                    carry_over_text = last_line + "\n"

    if not parsed_documents:
        if needs_ocr and not OCR_AVAILABLE:
            raise ValueError(
                "PDF 是扫描版，需要安装 OCR 依赖。请运行：pip install pytesseract pdf2image，并安装 Tesseract-OCR"
            )
        raise ValueError("PDF 中没有可读取的文字内容")
    return parsed_documents


def parse_word_file(raw_content: bytes, filename: str) -> list[ParsedDocument]:
    try:
        document = WordDocument(BytesIO(raw_content))
    except Exception as error:
        raise ValueError("Word 文件无法读取，请确认它是有效的 .docx 文件") from error

    blocks = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append(" | ".join(values))

    return [
        ParsedDocument(
            title=get_document_title(filename),
            content=normalize_text("\n".join(filter(None, blocks)), filename),
            source=f"上传 Word：{filename}",
        )
    ]


def parse_excel_file(raw_content: bytes, filename: str) -> list[ParsedDocument]:
    try:
        workbook = load_workbook(
            BytesIO(raw_content),
            read_only=True,
            data_only=True,
        )
    except Exception as error:
        raise ValueError("Excel 文件无法读取，请确认它是有效的 .xlsx 或 .xlsm 文件") from error

    base_title = get_document_title(filename)
    parsed_documents: list[ParsedDocument] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() if value is not None else "" for value in row]
                if any(values):
                    rows.append(" | ".join(values))
            if rows:
                parsed_documents.append(
                    ParsedDocument(
                        title=f"{base_title}（{sheet.title}）"[:200],
                        content=normalize_text("\n".join(rows), f"Excel 工作表 {sheet.title}"),
                        source=f"上传 Excel：{filename}（工作表：{sheet.title}）",
                    )
                )
    finally:
        workbook.close()

    if not parsed_documents:
        raise ValueError("Excel 中没有可读取的单元格内容")
    return parsed_documents


def parse_uploaded_file(filename: str, raw_content: bytes) -> list[ParsedDocument]:
    if len(raw_content) > MAX_DOCUMENT_BYTES:
        raise ValueError("单个文件不能超过 5MB")

    suffix = validate_upload_file(filename)
    if suffix in {".md", ".txt"}:
        return parse_text_file(raw_content, filename)
    if suffix == ".pdf":
        return parse_pdf_file(raw_content, filename)
    if suffix == ".docx":
        return parse_word_file(raw_content, filename)
    if suffix in {".xlsx", ".xlsm"}:
        return parse_excel_file(raw_content, filename)
    raise ValueError("无法根据文件类型选择解析器")


def fetch_web_page(url: str) -> tuple[str, str, str]:
    """Download a public HTML page and return title, text, and final URL."""
    parsed_url = urlparse(url)
    if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
        raise ValueError("网页链接必须是有效的 http:// 或 https:// 地址")

    request = Request(url, headers={"User-Agent": "MedicalHealthAssistant/1.0"})
    try:
        with urlopen(request, timeout=10) as response:
            content_type = response.headers.get_content_type()
            if content_type not in {"text/html", "application/xhtml+xml"}:
                raise ValueError("链接返回的不是网页 HTML 内容")
            raw_content = response.read(MAX_WEB_BYTES + 1)
            if len(raw_content) > MAX_WEB_BYTES:
                raise ValueError("网页内容不能超过 1MB")
            charset = response.headers.get_content_charset() or "utf-8"
            final_url = response.geturl()
    except HTTPError as error:
        raise ValueError(f"网页请求失败，状态码：{error.code}") from error
    except URLError as error:
        raise ValueError("网页无法访问，请检查链接或网络连接") from error
    except TimeoutError as error:
        raise ValueError("网页访问超时") from error

    try:
        html = raw_content.decode(charset, errors="replace")
    except LookupError:
        html = raw_content.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    for element in soup(["script", "style", "noscript", "svg", "nav", "footer", "header", "form"]):
        element.decompose()

    title = ""
    if soup.title and soup.title.string:
        title = soup.title.string.strip()
    text = soup.get_text("\n", strip=True)
    return (title[:200] or "网页资料", normalize_text(text, "网页"), final_url)


def parse_web_page(url: str, title: str | None = None) -> ParsedDocument:
    parsed_title, content, final_url = fetch_web_page(url)
    return ParsedDocument(
        title=(title or parsed_title).strip()[:200],
        content=content,
        source=f"网页导入：{final_url}",
        source_url=final_url,
    )
